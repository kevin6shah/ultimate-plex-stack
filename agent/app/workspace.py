from __future__ import annotations

import csv
import io
import json
import os
import re
import subprocess
from pathlib import Path


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


class Workspace:
    def __init__(self, root: str):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def resolve(self, relative_path: str) -> Path:
        candidate = (self.root / relative_path).resolve()
        if self.root not in candidate.parents and candidate != self.root:
            raise ValueError("path escapes workspace")
        return candidate

    def list_files(self) -> list[str]:
        visible: list[str] = []
        for path in self.root.rglob("*"):
            if not path.is_file():
                continue
            relative = path.relative_to(self.root)
            if relative.name == "claim.json":
                continue
            if any(part.startswith(".") for part in relative.parts):
                continue
            visible.append(str(relative))
        return sorted(visible)

    def read_text(self, relative_path: str, limit: int = 6000) -> str:
        path = self.resolve(relative_path)
        return path.read_text(encoding="utf-8", errors="replace")[:limit]

    def write_text(self, relative_path: str, content: str) -> str:
        path = self.resolve(relative_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return str(path.relative_to(self.root))

    def write_pdf(self, relative_path: str, title: str, body_text: str) -> str:
        path = self.resolve(relative_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        lines = [title.strip()] if title.strip() else []
        if lines:
            lines.append("")
        lines.extend((body_text or "").splitlines() or [""])
        wrapped: list[str] = []
        for raw_line in lines:
            line = raw_line.rstrip()
            if not line:
                wrapped.append("")
                continue
            while len(line) > 92:
                split_at = line.rfind(" ", 0, 92)
                if split_at <= 0:
                    split_at = 92
                wrapped.append(line[:split_at].rstrip())
                line = line[split_at:].lstrip()
            wrapped.append(line)

        lines_per_page = 48
        pages = [wrapped[index:index + lines_per_page] for index in range(0, len(wrapped), lines_per_page)] or [[""]]
        objects: list[bytes] = []

        def add_object(payload: str | bytes) -> int:
            data = payload.encode("utf-8") if isinstance(payload, str) else payload
            objects.append(data)
            return len(objects)

        font_id = add_object("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
        page_ids: list[int] = []
        content_ids: list[int] = []

        for page_lines in pages:
            stream_lines = ["BT", "/F1 11 Tf", "72 770 Td", "14 TL"]
            for line in page_lines:
                stream_lines.append(f"({_pdf_escape(line)}) Tj")
                stream_lines.append("T*")
            stream_lines.append("ET")
            stream = "\n".join(stream_lines).encode("utf-8")
            content_id = add_object(b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream")
            content_ids.append(content_id)
            page_ids.append(0)

        pages_id = add_object("<< /Type /Pages /Kids [] /Count 0 >>")
        for index, content_id in enumerate(content_ids):
            page_obj = (
                f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
            )
            page_ids[index] = add_object(page_obj)

        kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
        objects[pages_id - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("utf-8")
        title_value = _pdf_escape(title.strip() or "Friday Report")
        catalog_id = add_object(f"<< /Type /Catalog /Pages {pages_id} 0 R >>")
        info_id = add_object(f"<< /Title ({title_value}) /Producer (Friday Hands) >>")

        pdf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
        offsets = [0]
        for index, obj in enumerate(objects, start=1):
            offsets.append(len(pdf))
            pdf.extend(f"{index} 0 obj\n".encode("ascii"))
            pdf.extend(obj)
            pdf.extend(b"\nendobj\n")
        xref_offset = len(pdf)
        pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
        pdf.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
        pdf.extend(
            (
                f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R /Info {info_id} 0 R >>\n"
                f"startxref\n{xref_offset}\n%%EOF\n"
            ).encode("ascii")
        )
        path.write_bytes(bytes(pdf))
        return str(path.relative_to(self.root))

    def preview_table(self, relative_path: str, rows: int = 10) -> str:
        path = self.resolve(relative_path)
        suffix = path.suffix.lower()
        if suffix == ".csv":
            with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
                reader = csv.reader(handle)
                output = []
                for index, row in enumerate(reader):
                    output.append(" | ".join(row))
                    if index + 1 >= rows:
                        break
                return "\n".join(output)
        if suffix in {".json"}:
            data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
            rendered = json.dumps(data, indent=2, ensure_ascii=True)
            return rendered[:6000]
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                output = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    output.append(" | ".join("" if value is None else str(value) for value in row))
                    if index + 1 >= rows:
                        break
                return "\n".join(output)
            finally:
                workbook.close()
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            output = []
            for page in reader.pages[:rows]:
                output.append((page.extract_text() or "").strip())
            return "\n\n".join(output)[:6000]
        return self.read_text(relative_path)

    def run_shell(self, command: str, timeout_seconds: int = 60) -> str:
        blocked_patterns = (
            r"\bsudo\b",
            r"\bshutdown\b",
            r"\breboot\b",
            r"\bhalt\b",
            r"\bpoweroff\b",
            r"\bmkfs\b",
            r"\bmount\b",
            r"\bumount\b",
            r"rm\s+-rf\s+/",
            r":\(\)\s*\{",
            r"curl\s+[^|]+\|\s*sh",
            r"wget\s+[^|]+\|\s*sh",
            r">\s*/dev/sd",
        )
        if any(re.search(pattern, command) for pattern in blocked_patterns):
            return "Blocked shell command by Friday tool policy."
        result = subprocess.run(
            command,
            shell=True,
            cwd=self.root,
            capture_output=True,
            text=True,
            timeout=timeout_seconds,
            env={
                "PATH": os.environ.get("PATH", ""),
                "HOME": str(self.root),
                "PYTHONUNBUFFERED": "1",
            },
        )
        output = (result.stdout or "") + ("\n" + result.stderr if result.stderr else "")
        return output[:7000]

    def run_python(self, code: str, timeout_seconds: int = 60) -> str:
        result = subprocess.run(
            ["python3", "-c", code],
            cwd=self.root,
            capture_output=True,
            text=True,
            timeout=timeout_seconds,
            env={
                "PATH": os.environ.get("PATH", ""),
                "HOME": str(self.root),
                "PYTHONUNBUFFERED": "1",
            },
        )
        output = (result.stdout or "") + ("\n" + result.stderr if result.stderr else "")
        return output[:7000]

    def workspace_snapshot(self) -> str:
        files = self.list_files()
        buffer = io.StringIO()
        for file_name in files[:200]:
            buffer.write(file_name)
            buffer.write("\n")
        return buffer.getvalue().strip()[:4000]
